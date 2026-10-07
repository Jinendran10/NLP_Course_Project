"""Plain-text extraction from PDF, DOCX and text resumes."""

from __future__ import annotations

import io
from os import PathLike
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


def read_file(path: str | PathLike[str]) -> str:
    path = Path(path)
    return read_bytes(path.read_bytes(), path.name)


def read_bytes(data: bytes, filename: str) -> str:
    """Extract text from an in-memory file; the extension of ``filename`` picks the reader."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _read_pdf(data)
    elif ext == ".docx":
        text = _read_docx(data)
    elif ext in (".txt", ".md"):
        text = _decode(data)
    elif ext == ".doc":
        raise ValueError("legacy .doc files are not supported; save the resume as .docx or PDF")
    else:
        raise ValueError(
            f"unsupported file type {ext or '(no extension)'}; expected one of {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    if not text.strip():
        hint = " (scanned PDFs need OCR first)" if ext == ".pdf" else ""
        raise ValueError(f"no text could be extracted from {filename}{hint}")
    return text


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _read_pdf(data: bytes) -> str:
    try:
        import pymupdf
    except ImportError:
        pymupdf = None

    if pymupdf is not None:
        try:
            with pymupdf.open(stream=data, filetype="pdf") as doc:
                return "\n".join(page.get_text() for page in doc)
        except Exception as exc:  # PyMuPDF raises several unrelated types for bad files
            raise ValueError(f"could not read PDF: {exc}") from exc

    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("reading PDFs needs PyMuPDF or pypdf: pip install pymupdf") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ValueError(f"could not read PDF: {exc}") from exc


def _read_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("reading .docx files needs python-docx: pip install python-docx") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"could not read .docx file: {exc}") from exc

    parts = [p.text for p in document.paragraphs]
    # Many resume templates lay out skills and dates in tables.
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                text = cell.text.strip()
                if text and (not cells or cells[-1] != text):  # merged cells repeat their text
                    cells.append(text)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
